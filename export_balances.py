#!/usr/bin/env python3
"""
Export brut des balances comptables mensuelles Pennylane — toutes sociétés.

Une seule variable d'environnement à fournir : PENNYLANE_API_KEYS, contenant
la liste JSON de tous les jetons API (un par société). Le script :

  1. lit cette liste et boucle sur chaque jeton ;
  2. pour chaque jeton, identifie la société (endpoint /me) — inutile de
     tenir à jour une liste de noms de sociétés, tout vient de l'API ;
  3. récupère les exercices fiscaux (fiscal_years) pour connaître la plage
     de mois disponible ;
  4. récupère la balance comptable (trial_balance) mois par mois ;
  5. écrit toutes les lignes de toutes les sociétés dans un unique classeur
     Excel : Société | Mois | Trimestre | Exercice fiscal | Compte | Libellé | Solde

Ajouter une société = ajouter son jeton dans le secret PENNYLANE_API_KEYS.
Rien d'autre à modifier (ni secret supplémentaire, ni variable, ni workflow).

Format attendu de PENNYLANE_API_KEYS (liste JSON de chaînes) :
    ["cle_societe_1", "cle_societe_2", "cle_societe_3"]

Le "Solde" est le mouvement net du mois (Débit − Crédit) tel que renvoyé par
l'API pour la période period_start/period_end = ce mois-là — pas un solde
cumulé depuis le début de l'exercice.

Le "Trimestre" suit l'année civile du mois (2024T1 = janvier-mars 2024, etc.),
indépendamment de l'exercice fiscal. L'"Exercice fiscal" est déduit des dates
de l'exercice Pennylane qui couvre ce mois (ex. "2024" si calé sur l'année
civile, "2024-2025" s'il chevauche deux années civiles).

Endpoints Pennylane utilisés, uniquement ceux-ci :
    https://pennylane.readme.io/reference/gettrialbalance
    https://pennylane.readme.io/reference/company-fiscal-years
    https://pennylane.readme.io/changelog/v2-new-user-profile-endpoint  (/me,
        uniquement pour connaître le nom de la société)
"""
import datetime
import json
import os
import sys
from calendar import monthrange

import requests
from openpyxl import Workbook
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

BASE_URL = "https://app.pennylane.com/api/external/v2"
TIMEOUT = 30

OUTPUT_DIR = "reports"
HEADERS = ["Société", "Mois", "Trimestre", "Exercice fiscal", "Compte", "Libellé", "Solde"]
EUR_FMT = '#,##0.00 €;-#,##0.00 €;-'


# ---------------------------------------------------------------------------
# Accès API
# ---------------------------------------------------------------------------

def load_api_keys():
    raw = os.environ.get("PENNYLANE_API_KEYS")
    if not raw or not raw.strip():
        sys.exit(
            "Erreur : la variable d'environnement PENNYLANE_API_KEYS est absente ou vide.\n"
            "Elle doit contenir une liste JSON de jetons API, ex. :\n"
            '  ["cle_societe_1", "cle_societe_2"]'
        )
    try:
        keys = json.loads(raw)
    except json.JSONDecodeError as exc:
        sys.exit(
            f"Erreur : PENNYLANE_API_KEYS n'est pas un JSON valide ({exc}).\n"
            "Format attendu : [\"cle_societe_1\", \"cle_societe_2\"]"
        )
    if not isinstance(keys, list) or not all(isinstance(k, str) for k in keys):
        sys.exit("Erreur : PENNYLANE_API_KEYS doit être une liste JSON de chaînes de caractères.")
    keys = [k.strip() for k in keys if k.strip()]
    if not keys:
        sys.exit("Erreur : PENNYLANE_API_KEYS ne contient aucun jeton exploitable.")
    return keys


def paginated_get(api_key, path, params=None):
    """Parcourt toutes les pages d'un endpoint Pennylane (pagination par curseur)."""
    headers = {"Authorization": f"Bearer {api_key}"}
    items = []
    cursor = None
    while True:
        query = dict(params or {})
        query.setdefault("limit", 1000)
        if cursor:
            query["cursor"] = cursor
        resp = requests.get(f"{BASE_URL}/{path}", headers=headers, params=query, timeout=TIMEOUT)
        resp.raise_for_status()
        data = resp.json()
        items.extend(data.get("items", []))
        if not data.get("has_more") or not data.get("next_cursor"):
            break
        cursor = data["next_cursor"]
    return items


def fetch_company_info(api_key):
    headers = {"Authorization": f"Bearer {api_key}"}
    resp = requests.get(f"{BASE_URL}/me", headers=headers, timeout=TIMEOUT)
    resp.raise_for_status()
    return resp.json().get("company", {})


def fetch_fiscal_years(api_key):
    return paginated_get(api_key, "fiscal_years", {"sort": "start", "limit": 100})


def fetch_trial_balance(api_key, period_start, period_end):
    return paginated_get(
        api_key, "trial_balance",
        {"period_start": period_start, "period_end": period_end, "limit": 1000},
    )


# ---------------------------------------------------------------------------
# Plage de mois disponible, à partir des exercices fiscaux
# ---------------------------------------------------------------------------

def quarter_label(a_date):
    """'2024T1' pour janvier-mars 2024, '2024T2' pour avril-juin 2024, etc.
    (basé sur l'année civile du mois, pas sur l'exercice fiscal.)"""
    quarter = (a_date.month - 1) // 3 + 1
    return f"{a_date.year}T{quarter}"


def fiscal_year_label(a_date, fiscal_years):
    """Renvoie le libellé de l'exercice fiscal couvrant a_date, ex. '2024' si
    l'exercice est calé sur l'année civile, ou '2024-2025' s'il chevauche deux
    années civiles. Renvoie une chaîne vide si aucun exercice ne correspond."""
    for fy in fiscal_years:
        fy_start = datetime.date.fromisoformat(fy["start"])
        fy_finish = datetime.date.fromisoformat(fy["finish"])
        if fy_start <= a_date <= fy_finish:
            if fy_start.year == fy_finish.year:
                return str(fy_start.year)
            return f"{fy_start.year}-{fy_finish.year}"
    return ""


def month_range_from_fiscal_years(fiscal_years, today=None):
    """Renvoie (premier_jour, dernier_jour, 'YYYY-MM') pour chaque mois couvert
    par les exercices fiscaux, du plus ancien exercice au mois en cours (sans
    dépasser aujourd'hui)."""
    today = today or datetime.date.today()
    if not fiscal_years:
        return []

    starts = [datetime.date.fromisoformat(fy["start"]) for fy in fiscal_years]
    finishes = [datetime.date.fromisoformat(fy["finish"]) for fy in fiscal_years]
    range_start = min(starts)
    range_end = min(max(finishes), today)

    months = []
    cursor = range_start.replace(day=1)
    end_marker = range_end.replace(day=1)
    while cursor <= end_marker:
        last_day = monthrange(cursor.year, cursor.month)[1]
        month_end = min(cursor.replace(day=last_day), today)
        months.append((cursor, month_end, cursor.strftime("%Y-%m")))
        if cursor.month == 12:
            cursor = cursor.replace(year=cursor.year + 1, month=1)
        else:
            cursor = cursor.replace(month=cursor.month + 1)
    return months


# ---------------------------------------------------------------------------
# Une société
# ---------------------------------------------------------------------------

def export_company(api_key, index):
    """Renvoie la liste des lignes [Société, Mois, Trimestre, Exercice fiscal,
    Compte, Libellé, Solde] pour ce jeton API, ou lève une exception en cas
    d'échec."""
    company = fetch_company_info(api_key)
    company_name = company.get("name") or f"Société inconnue #{index}"
    print(f"[{index}] Société : {company_name} (ID {company.get('id')})")

    fiscal_years = fetch_fiscal_years(api_key)
    if not fiscal_years:
        print(f"[{index}]   Aucun exercice fiscal renvoyé, société ignorée.")
        return []

    months = month_range_from_fiscal_years(fiscal_years)
    if not months:
        print(f"[{index}]   Aucun mois à récupérer.")
        return []
    print(f"[{index}]   {len(months)} mois à récupérer, de {months[0][2]} à {months[-1][2]}.")

    rows = []
    for month_start, month_end, label in months:
        accounts = fetch_trial_balance(api_key, month_start.isoformat(), month_end.isoformat())
        trimestre = quarter_label(month_start)
        exercice = fiscal_year_label(month_start, fiscal_years)
        for a in accounts:
            solde = float(a.get("debits") or 0) - float(a.get("credits") or 0)
            rows.append([
                company_name,
                label,
                trimestre,
                exercice,
                a.get("formatted_number") or a.get("number"),
                a.get("label"),
                solde,
            ])
    print(f"[{index}]   {len(rows)} ligne(s) au total pour cette société.")
    return rows


# ---------------------------------------------------------------------------
# Classeur final
# ---------------------------------------------------------------------------

def build_workbook(rows):
    rows.sort(key=lambda r: (r[0], r[1], r[4]))  # Société, Mois, Compte

    wb = Workbook()
    ws = wb.active
    ws.title = "Balances brutes"

    ws.append(HEADERS)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in rows:
        ws.append(row)

    last_row = len(rows) + 1
    if last_row > 1:
        for r in range(2, last_row + 1):
            ws.cell(row=r, column=7).number_format = EUR_FMT
        table = Table(displayName="BalancesBrutes", ref=f"A1:G{last_row}")
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium9", showRowStripes=True)
        ws.add_table(table)

    for i, w in enumerate([28, 10, 10, 14, 14, 45, 16], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    return wb


# ---------------------------------------------------------------------------
# Programme principal
# ---------------------------------------------------------------------------

def main():
    api_keys = load_api_keys()
    print(f"{len(api_keys)} jeton(s) API à traiter.")

    all_rows = []
    failures = []
    for idx, api_key in enumerate(api_keys, start=1):
        try:
            all_rows.extend(export_company(api_key, idx))
        except requests.exceptions.RequestException as exc:
            detail = ""
            if getattr(exc, "response", None) is not None:
                detail = f" — réponse API : {exc.response.text[:300]}"
            print(f"[{idx}]   ÉCHEC : {exc}{detail}")
            failures.append(idx)

    if not all_rows:
        sys.exit("Erreur : aucune donnée récupérée pour aucune société, arrêt.")

    wb = build_workbook(all_rows)

    os.makedirs(OUTPUT_DIR, exist_ok=True)
    dated_path = os.path.join(OUTPUT_DIR, f"balances_{datetime.date.today().strftime('%Y-%m')}.xlsx")
    latest_path = os.path.join(OUTPUT_DIR, "balances_dernier.xlsx")
    wb.save(dated_path)
    wb.save(latest_path)

    print(f"\nClasseur écrit : {dated_path} et {latest_path} ({len(all_rows)} lignes au total).")
    if failures:
        print(f"Attention : {len(failures)} société(s) sur {len(api_keys)} ont échoué "
              f"(jetons n° {failures}) — voir les erreurs ci-dessus.")


if __name__ == "__main__":
    main()
